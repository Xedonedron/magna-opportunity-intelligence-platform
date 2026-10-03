"""Tests for Prospecting Hub & Lusha API Outbound Integration."""

from unittest.mock import AsyncMock, patch
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.company import Company
from app.models.company_contact import CompanyContact
from app.models.opportunity import Opportunity, TimelineEvent
from app.models.user import User


@pytest.fixture(autouse=True)
def enable_prospecting_in_tests(monkeypatch):
    from app.core.config import get_settings
    settings = get_settings()
    monkeypatch.setattr(settings, "ENABLE_PROSPECTING", True)


class TestProspectingDisabledGlobally:
    def test_prospecting_disabled_for_all(self, client, auth_headers, monkeypatch):
        from app.core.config import get_settings
        settings = get_settings()
        monkeypatch.setattr(settings, "ENABLE_PROSPECTING", False)

        resp = client.get("/api/prospecting/lusha/usage", headers=auth_headers)
        assert resp.status_code == 403
        assert "Fitur Prospecting dinonaktifkan" in resp.json()["detail"]


class TestProspectingLushaUsage:
    """Test Lusha Usage Endpoint."""

    @patch("app.api.prospecting.lusha_service.get_account_usage", new_callable=AsyncMock)
    def test_get_usage_success(self, mock_get_usage, client: TestClient):
        mock_get_usage.return_value = {
            "plan": "Starter",
            "credits_total": 4820,
            "credits_used": 1838,
            "credits_remaining": 2982,
            "features": {"contactSearch": 1, "revealEmail": 1, "revealPhone": 5},
            "rate_limit_per_minute": 40,
            "rate_limit_per_day": 100,
        }

        response = client.get("/api/prospecting/lusha/usage")
        assert response.status_code == 200
        data = response.json()
        assert data["credits_remaining"] == 2982
        assert data["plan"] == "Starter"


class TestProspectingLushaSearch:
    """Test Lusha Prospecting Search Endpoint."""

    def test_search_missing_company_name(self, client: TestClient):
        response = client.post("/api/prospecting/lusha/search", json={"company_name": ""})
        assert response.status_code in (400, 422)

    @patch("app.api.prospecting.lusha_service.search_contacts", new_callable=AsyncMock)
    def test_search_contacts_success(self, mock_search, client: TestClient):
        mock_search.return_value = {
            "success": True,
            "total": 1,
            "page": 0,
            "contacts": [
                {
                    "id": "123456",
                    "full_name": "Agus Pratama",
                    "first_name": "Agus",
                    "last_name": "Pratama",
                    "job_title": "Head of IT Infrastructure",
                    "company_name": "PT Bank Mega Tbk",
                    "company_domain": "bankmega.com",
                    "location": "Jakarta, Indonesia",
                    "linkedin_url": "https://linkedin.com/in/agus-pratama",
                    "has_email": True,
                    "has_phone": True,
                }
            ],
            "message": "Ditemukan 1 kontak terverifikasi di Lusha.",
        }

        response = client.post(
            "/api/prospecting/lusha/search",
            json={
                "company_name": "Bank Mega",
                "country": "Indonesia",
                "seniority": "head_lead",
                "job_function": "network",
            },
        )
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert len(data["contacts"]) == 1
        assert data["contacts"][0]["full_name"] == "Agus Pratama"

    @patch("app.api.prospecting.lusha_service.search_contacts", new_callable=AsyncMock)
    def test_search_contacts_with_domain_and_titles(self, mock_search, client: TestClient):
        mock_search.return_value = {
            "success": True,
            "total": 2,
            "page": 0,
            "contacts": [
                {
                    "id": "c1",
                    "full_name": "Budi Santoso",
                    "first_name": "Budi",
                    "last_name": "Santoso",
                    "job_title": "Chief Information Officer",
                    "company_name": "PT Bank Mega Tbk",
                    "company_domain": "bankmega.com",
                    "has_email": True,
                    "has_phone": True,
                },
                {
                    "id": "c2",
                    "full_name": "Siti Rahma",
                    "first_name": "Siti",
                    "last_name": "Rahma",
                    "job_title": "Head of IT Security",
                    "company_name": "PT Bank Mega Tbk",
                    "company_domain": "bankmega.com",
                    "has_email": True,
                    "has_phone": False,
                },
            ],
            "message": "Ditemukan 2 kontak terverifikasi di Lusha.",
        }

        response = client.post(
            "/api/prospecting/lusha/search",
            json={
                "company_name": "PT Bank Mega Tbk",
                "company_domain": "bankmega.com",
                "job_titles": ["Chief Information Officer", "Head of IT Security"],
                "limit": 25,
            },
        )
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert len(data["contacts"]) == 2
        assert data["contacts"][0]["job_title"] == "Chief Information Officer"
        mock_search.assert_called_once_with(
            company_name="PT Bank Mega Tbk",
            company_domain="bankmega.com",
            country="Indonesia",
            seniority=None,
            job_function=None,
            job_titles=["Chief Information Officer", "Head of IT Security"],
            page=0,
            limit=25,
        )


class TestProspectingLushaEnrich:
    """Test Lusha Contact Enrichment Endpoint."""

    @patch("app.api.prospecting.lusha_service.enrich_contact", new_callable=AsyncMock)
    def test_enrich_contact_success(self, mock_enrich, client: TestClient):
        mock_enrich.return_value = {
            "success": True,
            "contact_id": "123456",
            "emails": ["agus.pratama@bankmega.com"],
            "phones": ["+6281234567890"],
            "message": "Kontak berhasil diperkaya dengan data terverifikasi Lusha.",
        }

        response = client.post(
            "/api/prospecting/lusha/enrich",
            json={"contact_id": "123456", "reveal": ["emails", "phones"]},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert "agus.pratama@bankmega.com" in data["emails"]
        assert len(data["phones"]) > 0

    @patch("app.api.prospecting.lusha_service.enrich_contact", new_callable=AsyncMock)
    def test_enrich_contact_singular_reveal_normalized(self, mock_enrich, client: TestClient):
        mock_enrich.return_value = {
            "success": True,
            "contact_id": "123456",
            "emails": ["agus.pratama@bankmega.com"],
            "phones": [],
            "email": "agus.pratama@bankmega.com",
            "phone": None,
            "message": "Kontak berhasil diperkaya dengan data terverifikasi Lusha.",
        }

        # Send singular 'email' and 'phone'
        response = client.post(
            "/api/prospecting/lusha/enrich",
            json={"contact_id": "123456", "reveal": ["email", "phone"]},
        )
        assert response.status_code == 200
        mock_enrich.assert_called_once()
        called_args = mock_enrich.call_args[1]
        assert called_args["reveal"] == ["emails", "phones"]
        data = response.json()
        assert data["success"] is True
        assert data["email"] == "agus.pratama@bankmega.com"


class TestProspectingGenerate:
    """Test Sales Hypothesis & Outreach Generation."""

    def test_generate_dossier_network(self, client: TestClient):
        payload = {
            "company_name": "PT Bank Mega Tbk",
            "industry": "Financial Services",
            "job_function": "network",
            "seniority": "head_lead",
            "contact": {
                "full_name": "Agus Pratama",
                "job_title": "Head of IT Network & Telecom",
                "email": "agus.pratama@bankmega.com",
                "phone": "+6281234567890",
            },
        }
        response = client.post("/api/prospecting/generate", json=payload)
        assert response.status_code == 200
        data = response.json()
        assert data["company_name"] == "PT Bank Mega Tbk"
        assert "SD-WAN" in data["hypothesis"]["recommended_solution"]
        assert len(data["hypothesis"]["pain_points"]) >= 3
        assert "Dan dari tim Solution Architect" in data["outreach"]["whatsapp"]
        assert len(data["discovery_questions"]) >= 4

    def test_generate_dossier_security_c_level(self, client: TestClient):
        payload = {
            "company_name": "PT Telkom Indonesia Tbk",
            "job_function": "security",
            "seniority": "c_level",
            "contact": {
                "full_name": "Budi Santoso",
                "job_title": "Chief Information Security Officer",
            },
        }
        response = client.post("/api/prospecting/generate", json=payload)
        assert response.status_code == 200
        data = response.json()
        assert "Zero-Trust" in data["hypothesis"]["recommended_solution"]
        assert "Solution Advisory" in data["outreach"]["whatsapp"]


class TestProspectingConvert:
    """Test Atomic Conversion of Prospect to Opportunity & Database Records."""

    def test_convert_to_opportunity_success(
        self,
        client: TestClient,
        lgo_auth_headers: dict[str, str],
        db: Session,
        lgo_user: User,
    ):
        payload = {
            "company_name": "PT Bank Mega Tbk",
            "industry": "Banking & Financial Services",
            "website": "https://www.bankmega.com",
            "contact": {
                "full_name": "Agus Pratama",
                "job_title": "Head of IT Network",
                "department": "IT",
                "email": "agus.pratama@bankmega.com",
                "phone": "+628****7890",
                "linkedin_url": "https://linkedin.com/in/agus-pratama",
            },
            "hypothesis": {
                "pain_points": [
                    "Kompleksitas routing multi-cabang & latensi tinggi",
                    "OPEX leased line konvensional tinggi",
                ],
                "recommended_solution": "Automated SD-WAN & Enterprise Campus Network",
                "solution_pillar": "Network & Enterprise Workplace",
                "tech_stack": ["Cisco Catalyst", "Fortinet Secure SD-WAN"],
                "key_benefits": ["Reduksi biaya bandwidth hingga 40%"],
            },
            "opportunity": {
                "title": "Modernisasi SD-WAN 300+ Cabang Bank Mega",
                "notes": "Target implementasi Q4 2026",
                "estimated_value": 750000000.0,
            },
        }

        response = client.post(
            "/api/prospecting/convert",
            json=payload,
            headers=lgo_auth_headers,
        )
        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "success"
        assert "opportunity_id" in data
        assert "company_id" in data
        assert "contact_id" in data

        # Verify Database state
        import uuid as _uuid
        company = db.query(Company).filter(Company.id == _uuid.UUID(data["company_id"])).first()
        assert company is not None
        assert "bank mega" in company.normalized_name

        contact = db.query(CompanyContact).filter(CompanyContact.id == _uuid.UUID(data["contact_id"])).first()
        assert contact is not None
        assert contact.email == "agus.pratama@bankmega.com"
        assert contact.company_id == company.id

        opp = db.query(Opportunity).filter(Opportunity.id == _uuid.UUID(data["opportunity_id"])).first()
        assert opp is not None
        assert opp.company_id == company.id
        assert opp.primary_contact_id == contact.id
        assert opp.potential_revenue == 750000000.0
        assert "SD-WAN" in opp.product

        timeline = (
            db.query(TimelineEvent)
            .filter(TimelineEvent.opportunity_id == opp.id)
            .first()
        )
        assert timeline is not None
        assert timeline.action == "Prospecting Converted"

    def test_convert_to_opportunity_forbidden_for_unauthorized_role(
        self,
        client: TestClient,
        auth_headers: dict[str, str],
    ):
        """Verify role without LGO/Manager/Superadmin or prospecting capability gets 403 Forbidden."""
        payload = {
            "company_name": "Unauthorized Access PT",
            "contact": {
                "full_name": "John Doe",
                "job_title": "Staff",
            },
        }
        response = client.post(
            "/api/prospecting/convert",
            json=payload,
            headers=auth_headers,
        )
        assert response.status_code == 403
        detail = response.json().get("detail", "")
        assert "Akses fitur Lusha Prospecting terbatas" in detail

    def test_prospecting_allowed_with_explicit_capability(
        self,
        client: TestClient,
        db: Session,
    ):
        """Verify a user with non-standard role can access prospecting if capability is granted."""
        import uuid as _uuid
        from app.services.auth import create_access_token

        custom_user = User(
            id=_uuid.uuid4(),
            email="specialist@smartnet.co.id",
            full_name="Specialist User",
            role="staff",
            capabilities="view,prospecting",
            is_active=True,
        )
        db.add(custom_user)
        db.commit()

        token = create_access_token(data={"sub": str(custom_user.id), "email": custom_user.email})
        headers = {"Authorization": f"Bearer {token}"}

        payload = {
            "company_name": "PT Fintek Nusantara",
            "industry": "Fintech",
            "job_function": "it",
            "seniority": "c_level",
            "contact": {
                "full_name": "Budi Santoso",
                "job_title": "CTO",
            },
        }
        response = client.post(
            "/api/prospecting/generate",
            json=payload,
            headers=headers,
        )
        assert response.status_code == 200
        data = response.json()
        assert data["company_name"] == "PT Fintek Nusantara"


class TestProspectingInteractiveFlow:
    """Test Company Disambiguation, Selective Enrich, Stakeholder Sync & Excel Export."""

    @patch("app.api.prospecting.lusha_service.search_companies", new_callable=AsyncMock)
    def test_search_companies_disambiguation(self, mock_search_comp, client: TestClient, db: Session):
        import uuid as _uuid
        # Prepopulate a local company
        comp = Company(
            id=_uuid.uuid4(),
            name="Bank OCBC NISP",
            normalized_name="bank ocbc nisp",
            root_domain="ocbcnisp.com",
            industry="Banking",
        )
        db.add(comp)
        db.commit()

        # Mock Lusha response
        mock_search_comp.return_value = [
            {
                "name": "OCBC Bank Singapore",
                "domain": "ocbc.com",
                "industry": "Financial Services",
                "country": "Singapore",
                "city": "Singapore",
                "employee_count": "10000+",
            }
        ]

        response = client.get("/api/prospecting/companies/search?q=OCBC")
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert len(data["results"]) >= 2
        # Check that local company has in_database = True
        local_item = next(r for r in data["results"] if r["name"] == "Bank OCBC NISP")
        assert local_item["in_database"] is True
        # Check that Lusha candidate has in_database = False
        lusha_item = next(r for r in data["results"] if r["name"] == "OCBC Bank Singapore")
        assert lusha_item["in_database"] is False

    @patch("app.api.prospecting.lusha_service.enrich_contacts", new_callable=AsyncMock)
    def test_enrich_contacts_selective(self, mock_enrich, client: TestClient):
        mock_enrich.return_value = [
            {
                "id": "c1",
                "fullName": "Jane Doe",
                "jobTitle": "Chief Digital Officer",
                "emails": [{"email": "jane@ocbc.com"}],
                "phones": [{"number": "+628****6789"}],
                "linkedinUrl": "https://linkedin.com/in/janedoe",
                "credits_charged": 1,
            }
        ]

        response = client.post(
            "/api/prospecting/lusha/enrich",
            json={
                "contact_ids": ["c1"],
                "reveal": ["emails"],
            },
        )
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert len(data["contacts"]) == 1
        assert data["contacts"][0]["emails"] == ["jane@ocbc.com"]
        assert data["credits_charged"] == 1

    def test_save_to_stakeholders_success(self, client: TestClient, db: Session):
        payload = {
            "company_name": "PT Astra Digital",
            "company_domain": "astradigital.id",
            "industry": "Technology",
            "country": "Indonesia",
            "contacts": [
                {
                    "name": "Budi Setiawan",
                    "job_title": "VP of Engineering",
                    "email": "budi@astradigital.id",
                    "phone": "+628111222333",
                }
            ],
        }
        response = client.post("/api/prospecting/save-to-stakeholders", json=payload)
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["saved_count"] == 1

        # Check DB
        contact = (
            db.query(CompanyContact)
            .filter(CompanyContact.email == "budi@astradigital.id")
            .first()
        )
        assert contact is not None
        assert contact.name == "Budi Setiawan"
        assert contact.job_title == "VP of Engineering"

    def test_export_excel_success(self, client: TestClient):
        import io
        import openpyxl

        payload = {
            "company_name": "PT Telkom Indonesia",
            "contacts": [
                {
                    "name": "Rudiantara",
                    "job_title": "Commissioner",
                    "email": "rudi@telkom.co.id",
                    "phone": "+62812999999",
                },
                {
                    "name": "Ririek Adriansyah",
                    "job_title": "President Director",
                    "email": "ririek@telkom.co.id",
                    "phone": None,
                },
            ],
        }
        response = client.post("/api/prospecting/export-excel", json=payload)
        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment; filename=" in response.headers.get("content-disposition", "")

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        assert ws is not None
        assert ws.title == "Stakeholders"
        # Row 4 should be header
        headers = [ws.cell(row=4, column=i).value for i in range(1, 6)]
        assert headers == ["No", "Nama", "Job Title / Jabatan", "Email", "Nomor Telepon"]
        # Row 5: first contact
        assert ws.cell(row=5, column=2).value == "Rudiantara"
        assert ws.cell(row=5, column=4).value == "rudi@telkom.co.id"
        # Row 6: second contact with missing phone falling back to "-"
        assert ws.cell(row=6, column=2).value == "Ririek Adriansyah"
        assert ws.cell(row=6, column=5).value == "-"

    def test_company_stakeholders_export_excel(
        self,
        client: TestClient,
        db: Session,
        auth_headers: dict[str, str],
    ):
        import io
        import uuid as _uuid
        import openpyxl

        comp = Company(
            id=_uuid.uuid4(),
            name="PT Indosat Ooredoo Hutchison",
            normalized_name="pt indosat ooredoo hutchison",
            root_domain="ioh.co.id",
            industry="Telecommunication",
        )
        db.add(comp)
        db.commit()

        c1 = CompanyContact(
            company_id=comp.id,
            name="Vikram Sinha",
            job_title="President Director & CEO",
            email="vikram@ioh.co.id",
            phone="+628****0001",
        )
        db.add(c1)
        db.commit()

        response = client.get(
            f"/api/companies/{comp.id}/contacts/export-excel",
            headers=auth_headers,
        )
        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        assert ws is not None
        assert ws.cell(row=5, column=2).value == "Vikram Sinha"
        assert ws.cell(row=5, column=4).value == "vikram@ioh.co.id"

