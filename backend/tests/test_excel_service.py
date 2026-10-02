"""Tests for Excel Export Service."""

import io
import openpyxl
import pytest

from app.services.excel_service import generate_contacts_excel


def test_generate_contacts_excel_success():
    company_name = "PT Bank OCBC NISP Tbk"
    contacts = [
        {
            "name": "Budi Santoso",
            "job_title": "Head of IT Infrastructure",
            "email": "budi.santoso@ocbcnisp.com",
            "phone": "+6281234567890",
        },
        {
            "name": "Siti Rahma",
            "job_title": "Enterprise Architect",
            "email": "siti.rahma@ocbcnisp.com",
            "phone": "+6281987654321",
        },
    ]

    excel_bytes = generate_contacts_excel(company_name, contacts)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Load and verify Excel contents
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws is not None

    # Title check
    assert company_name in ws["A1"].value

    # Header row check (Row 4)
    expected_headers = ["No", "Nama", "Job Title / Jabatan", "Email", "Nomor Telepon"]
    row_4_values = [ws.cell(row=4, column=col).value for col in range(1, 6)]
    assert row_4_values == expected_headers

    # Row 5 (Contact 1)
    assert ws.cell(row=5, column=1).value == 1
    assert ws.cell(row=5, column=2).value == "Budi Santoso"
    assert ws.cell(row=5, column=3).value == "Head of IT Infrastructure"
    assert ws.cell(row=5, column=4).value == "budi.santoso@ocbcnisp.com"
    assert ws.cell(row=5, column=5).value == "+6281234567890"

    # Row 6 (Contact 2)
    assert ws.cell(row=6, column=1).value == 2
    assert ws.cell(row=6, column=2).value == "Siti Rahma"
    assert ws.cell(row=6, column=3).value == "Enterprise Architect"
    assert ws.cell(row=6, column=4).value == "siti.rahma@ocbcnisp.com"
    assert ws.cell(row=6, column=5).value == "+6281987654321"


def test_generate_contacts_excel_empty_and_fallback():
    company_name = "OCBC"
    contacts = [
        {
            "full_name": "John Doe",  # test fallback from full_name
            "position": "Director",   # test fallback from position
            # email missing
            # phone missing
        }
    ]

    excel_bytes = generate_contacts_excel(company_name, contacts)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    assert ws.cell(row=5, column=1).value == 1
    assert ws.cell(row=5, column=2).value == "John Doe"
    assert ws.cell(row=5, column=3).value == "Director"
    assert ws.cell(row=5, column=4).value == "-"
    assert ws.cell(row=5, column=5).value == "-"
